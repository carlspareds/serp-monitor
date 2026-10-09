"""Last small fixes.

1. Every report route prints the same "Project › Subject" label (PDF text,
   Excel title, report sidecar): Reports dialog, legacy /export routes,
   weekly job and the assistant.
2. The Apify cost estimate never says "about $0.00".
3. Assistant list_results (read, viewer allowed) + its results card.
4. URL identity ignores known locale path prefixes (Spotify /intl-xx, Apple /xx).
5. Times shown with "UTC".
6. Eval claim_note: opening Notes OR a backed save_note proposal.
No real network: Apify, AI and email are stubbed or unconfigured.
"""
import base64
import importlib.util
import json
import re
import sqlite3
import zlib
from pathlib import Path

import openpyxl
import pytest
from conftest import _ROOT

from assistant import cards, tools
from database.models import Database, normalize_url, url_hash

LABEL = "Demo › Jordan Ashworth (demo)"
FRAG = {"X-SM-Fragment": "1"}


def _demo_names(db, subject_id: int = 1) -> None:
    """Project 1 = "Demo", subject 1 = "Jordan Ashworth (demo)"."""
    db.rename_client(1, "Demo")
    db.update_subject(subject_id, name="Jordan Ashworth (demo)")


# --------------------------------------------------------------------- #
# PDF text (no PDF library in requirements: content streams, Tj strings)
# --------------------------------------------------------------------- #
_STREAM = re.compile(rb"stream\r?\n(.*?)\s*endstream", re.DOTALL)
_STRING = re.compile(rb"\((?:\\.|[^\\)])*\)", re.DOTALL)
_ESC = {b"n": b"\n", b"r": b"\r", b"t": b"\t", b"b": b"\b", b"f": b"\f",
        b"(": b"(", b")": b")", b"\\": b"\\"}


def _unescape(raw: bytes) -> bytes:
    out, i = bytearray(), 0
    while i < len(raw):
        c = raw[i:i + 1]
        if c != b"\\":
            out += c
            i += 1
            continue
        nxt = raw[i + 1:i + 2]
        octal = re.match(rb"[0-7]{1,3}", raw[i + 1:i + 4])
        if octal:
            out.append(int(octal.group(), 8) & 0xFF)
            i += 1 + len(octal.group())
        else:
            out += _ESC.get(nxt, nxt)
            i += 2
    return bytes(out)


def pdf_text(path) -> str:
    """Text drawn on the PDF's pages (Helvetica, cp1252), lines joined by newlines."""
    data = Path(path).read_bytes()
    lines = []
    for m in _STREAM.finditer(data):
        body = m.group(1)
        if body.rstrip().endswith(b"~>"):        # ReportLab: ASCII85 then Flate
            try:
                body = base64.a85decode(body.rstrip()[:-2].replace(b"\n", b""))
            except ValueError:
                pass
        try:
            body = zlib.decompress(body)
        except zlib.error:
            pass
        if b"Tj" not in body and b"TJ" not in body:
            continue
        for s in _STRING.finditer(body):
            lines.append(_unescape(s.group()[1:-1]).decode("cp1252", "replace"))
    return "\n".join(lines)


def test_pdf_text_helper_reads_a_built_pdf(seeded_db, tmp_path):
    from export.pdf_report import build_pdf
    out = tmp_path / "r.pdf"
    build_pdf(seeded_db, str(out), subject_name="Acme ▸ Ann (x)")
    assert "Acme › Ann (x) · Executive report" in pdf_text(out)


# --------------------------------------------------------------------- #
# 1. One label on every report route
# --------------------------------------------------------------------- #
@pytest.fixture
def demo_client(client, seeded_db, tmp_path, monkeypatch):
    from dashboard import deps
    _demo_names(seeded_db)
    monkeypatch.setattr(deps, "exports_dir", tmp_path / "exports")
    return client


def _sidecar(base: Path, name: str) -> dict:
    return json.loads((base / "1" / f"{name}.json").read_text())


def test_reports_dialog_pdf_has_project_and_subject(demo_client, tmp_path):
    r = demo_client.post("/reports/create", data={"kind": "pdf", "subject_id": "1"}, headers=FRAG)
    assert r.status_code == 200, r.text[:300]
    base = tmp_path / "exports"
    pdf = next((base / "1").glob("*.pdf"))
    assert LABEL in pdf_text(pdf)
    assert _sidecar(base, pdf.name)["subject"] == LABEL


@pytest.mark.parametrize("kind", ["xlsx", "csv"])
def test_reports_dialog_workbook_and_csv_carry_the_label(demo_client, tmp_path, kind):
    r = demo_client.post("/reports/create", data={"kind": kind, "subject_id": "1"}, headers=FRAG)
    assert r.status_code == 200
    base = tmp_path / "exports"
    f = next((base / "1").glob(f"*.{kind}"))
    assert _sidecar(base, f.name)["subject"] == LABEL
    if kind == "xlsx":
        assert openpyxl.load_workbook(f)["SUM"]["A1"].value == f"SERP MONITOR  ·  {LABEL}"


def test_legacy_export_routes_use_the_same_label(demo_client, tmp_path):
    r = demo_client.get("/export/pdf")
    assert r.status_code == 200
    pdf = tmp_path / "pdf.pdf"
    pdf.write_bytes(r.content)
    assert LABEL in pdf_text(pdf)
    r = demo_client.post("/export/run", data={})
    assert r.status_code == 200
    book = tmp_path / "b.xlsx"
    book.write_bytes(r.content)
    assert openpyxl.load_workbook(book)["SUM"]["A1"].value == f"SERP MONITOR  ·  {LABEL}"


def test_assistant_report_uses_the_same_label(seeded_db, tmp_path):
    _demo_names(seeded_db)
    ctx = tools.ToolContext(db=seeded_db.for_subject(1), subject_id=1, role="operator",
                            user="op", exports_dir=tmp_path / "exports")
    out = tools.exec_create_report(ctx, {"kind": "pdf", "period": "30d"})
    pdf = tmp_path / "exports" / "1" / out["file"]
    assert LABEL in pdf_text(pdf)
    out = tools.exec_create_report(ctx, {"kind": "csv", "period": "all"})
    assert _sidecar(tmp_path / "exports", out["file"])["subject"] == LABEL


def test_weekly_job_workbook_uses_the_same_label(tmp_path, monkeypatch):
    from test_weekly import _FakeRunner

    from report import weekly
    db = Database(db_path=str(tmp_path / "w.db"))
    _demo_names(db)
    monkeypatch.setattr(weekly, "project_root", tmp_path)
    monkeypatch.setattr(weekly, "ApifyRunner", lambda: _FakeRunner())
    monkeypatch.setattr(weekly.email_sender, "is_configured", lambda: False)
    weekly.run(db=db.for_subject(1), query="Jordan Ashworth", countries=["us"])
    base = tmp_path / "data" / "exports"
    book = next((base / "1").glob("weekly_*.xlsx"))
    assert openpyxl.load_workbook(book)["SUM"]["A1"].value == f"SERP MONITOR  ·  {LABEL}"
    assert _sidecar(base, book.name)["subject"] == LABEL


def test_create_report_ignores_a_name_only_label(seeded_db, tmp_path, monkeypatch):
    """The Reports dialog used to pass the name only; create_report now
    always builds the title itself."""
    from export import pdf_report, reports
    _demo_names(seeded_db)
    seen = {}
    real = pdf_report.build_pdf
    monkeypatch.setattr(pdf_report, "build_pdf", lambda *a, **k: (seen.update(k), real(*a, **k))[1])
    reports.create_report(seeded_db.for_subject(1), tmp_path, 1, "pdf")
    assert seen["subject_name"] == LABEL


# --------------------------------------------------------------------- #
# 2. Cost estimate
# --------------------------------------------------------------------- #
def test_cost_phrase_never_says_zero():
    from common.money import run_cost_phrase
    assert run_cost_phrase(0.0004) == "under $0.01 per country and engine"
    assert run_cost_phrase(0) == "under $0.01 per country and engine"
    assert run_cost_phrase(0.0004, runs=6) == "under $0.01 per country and engine"
    assert run_cost_phrase(0.03) == "about $0.03"
    assert run_cost_phrase(0.03, runs=4) == "about $0.12"


def _cheap_run(db):
    db.save_run(apify_run_id="cheap", queries=["q"], country_code="gb", status="completed",
                summary={"total": 0, "positive": 0, "negative": 0, "neutral": 0}, cost_usd=0.0007)


def test_new_search_dialog_estimate_for_cheap_runs(client, seeded_db):
    _cheap_run(seeded_db)
    html = client.get("/ui/search", headers=FRAG).text
    assert "under $0.01 per country and engine" in html and "$0.00" not in html


def test_assistant_search_card_estimate_for_cheap_runs(seeded_db, monkeypatch):
    monkeypatch.setenv("APIFY_API_KEY", "apify_api_test")
    seeded_db.set_tracked_subject(query="Jordan Ashworth")
    _cheap_run(seeded_db)
    ctx = tools.ToolContext(db=seeded_db.for_subject(1), subject_id=1, role="operator")
    card = tools.prepare_start_search(ctx, {"countries": ["us", "gb"]})
    assert "under $0.01 per country and engine" in card["cost_note"]
    assert "$0.00" not in card["cost_note"]


def test_assistant_card_js_mirrors_the_cost_wording():
    js = (_ROOT / "static" / "app.js").read_text()
    assert "'under $0.01 per country and engine'" in js


# --------------------------------------------------------------------- #
# 3. list_results
# --------------------------------------------------------------------- #
def _viewer(db):
    return tools.ToolContext(db=db.for_subject(1), subject_id=1, role="viewer")


def test_list_results_is_a_registered_viewer_read_tool():
    t = tools.get_tool("list_results")
    assert t.kind == "read" and t.min_role == "viewer"
    assert t.parameters["properties"]["limit"]["maximum"] == 50
    assert "list_results" in (_ROOT / "docs" / "assistant-knowledge.md").read_text() \
        or "latest run's results" in (_ROOT / "docs" / "assistant-knowledge.md").read_text()
    assert "`list_results(" in (_ROOT / "docs" / "assistant.md").read_text()


def test_list_results_returns_the_latest_run_by_label(seeded_db):
    t = tools.get_tool("list_results")
    out = t.executor(_viewer(seeded_db), tools.validate_args(t, {}))
    assert out["total"] == 5 and out["count_by_label"] == {"positive": 2, "negative": 2, "neutral": 1}
    first = out["items"][0]
    assert set(first) >= {"position", "title", "domain", "label", "source", "same_person"}
    assert [i["position"] for i in out["items"]] == sorted(i["position"] for i in out["items"])
    assert {i["source"] for i in out["items"]} <= {"human", "rule", "ai"}
    neg = t.executor(_viewer(seeded_db), tools.validate_args(t, {"label": "negative", "limit": 99}))
    assert neg["total"] == 2 and all(i["label"] == "negative" for i in neg["items"])
    one = t.executor(_viewer(seeded_db), tools.validate_args(t, {"limit": 1, "country": "us"}))
    assert len(one["items"]) == 1 and one["more"] == 4
    with pytest.raises(tools.ToolError, match="No run for GB"):
        t.executor(_viewer(seeded_db), tools.validate_args(t, {"country": "gb"}))


def test_list_results_marks_human_labels(seeded_db):
    row = next(r for r in seeded_db.get_results_filtered(limit=50) if r["domain"] == "scam.example")
    seeded_db.set_review_verdict(row["id"], "neutral")
    t = tools.get_tool("list_results")
    out = t.executor(_viewer(seeded_db), {})
    item = next(i for i in out["items"] if i["domain"] == "scam.example")
    assert item["source"] == "human" and item["label"] == "neutral"


def test_list_results_card_for_a_viewer(seeded_db):
    from assistant import harness
    from assistant.fakes import ScriptedModel, call, say
    from assistant.pending import PendingStore
    model = ScriptedModel([call("list_results", label="negative"), say("Two negative results.")])
    res = harness.run_turn(model, _viewer(seeded_db),
                           [{"role": "user", "content": "Table of current results by label"}],
                           session_key="s1", store=PendingStore())
    card = next(e for e in res.events if e.get("type") == "card")
    assert card["kind"] == "results" and card["tool"] == "list_results"
    assert card["data"]["can_review"] is False and len(card["data"]["items"]) == 2
    assert card["data"]["items"][0]["position"] is not None


def test_results_card_builder_caps_rows(seeded_db):
    t = tools.get_tool("list_results")
    out = t.executor(_viewer(seeded_db), {})
    card = cards.from_tool("list_results", out, _viewer(seeded_db))
    assert card["kind"] == "results" and card["data"]["title"].startswith("Latest run · 5 results")
    assert len(card["data"]["items"]) == 5 and card["data"]["more"] == 0


# --------------------------------------------------------------------- #
# 4. Locale path prefixes
# --------------------------------------------------------------------- #
@pytest.mark.parametrize("a,b", [
    ("https://open.spotify.com/intl-es/artist/3TVXtAsR1Inumwj472S9r4",
     "https://open.spotify.com/artist/3TVXtAsR1Inumwj472S9r4"),
    ("https://open.spotify.com/intl-pt-BR/track/abc?si=1", "https://open.spotify.com/track/abc?si=1"),
    ("https://music.apple.com/us/album/x/123", "https://music.apple.com/gb/album/x/123"),
    ("https://podcasts.apple.com/es/podcast/show/id1", "https://podcasts.apple.com/podcast/show/id1"),
])
def test_locale_prefixes_are_the_same_page(a, b):
    assert normalize_url(a) == normalize_url(b) and url_hash(a) == url_hash(b)


@pytest.mark.parametrize("a,b", [
    ("https://example.com/intl-es/artist/1", "https://example.com/artist/1"),      # other host
    ("https://open.spotify.com/international/x", "https://open.spotify.com/x"),     # not intl-xx
    ("https://music.apple.com/us/browse", "https://music.apple.com/browse"),        # not content
    ("https://www.apple.com/es/iphone/", "https://www.apple.com/iphone/"),          # not music
])
def test_other_paths_are_left_alone(a, b):
    assert normalize_url(a) != normalize_url(b)


def test_url_key_version_bump_recomputes_keys(tmp_path):
    from database import base
    assert base.URL_KEY_VERSION == "3"
    path = str(tmp_path / "k.db")
    db = Database(db_path=path)
    db.save_run(apify_run_id="r1", queries=["q"], country_code="us", status="completed",
                summary={"total": 1, "positive": 0, "negative": 0, "neutral": 1})
    db.save_results([{"query": "q", "url": "https://open.spotify.com/intl-de/artist/z", "title": "t",
                      "description": "", "position": 1, "country_code": "us",
                      "domain": "open.spotify.com"}], run_id="r1")
    conn = sqlite3.connect(path)
    conn.execute("UPDATE search_results SET url_key = 'stale'")
    conn.execute("UPDATE settings SET value = '2' WHERE key = 'url_key_version'")
    conn.commit()
    conn.close()
    Database(db_path=path)
    conn = sqlite3.connect(path)
    key = conn.execute("SELECT url_key FROM search_results").fetchone()[0]
    conn.close()
    assert key == url_hash("https://open.spotify.com/artist/z")


def test_repair_script_mentions_locale_prefixes():
    text = (_ROOT / "scripts" / "repair_url_identity.py").read_text()
    assert "locale" in text and "URL_KEY_VERSION" in text


# --------------------------------------------------------------------- #
# 5. UTC
# --------------------------------------------------------------------- #
def test_integration_tested_time_says_utc():
    from dashboard import integrations as integ
    assert integ.working_text("2026-10-09T02:34:00") == "Working · tested 9 Oct 02:34 UTC"
    assert integ.working_text("2026-10-09") == "Working · tested 9 Oct"


def test_assistant_prompt_asks_for_utc_times():
    from assistant import prompts
    assert "UTC" in prompts.RULES


# --------------------------------------------------------------------- #
# 6. Eval: claim_note accepts Notes or a backed save_note proposal
# --------------------------------------------------------------------- #
def _eval_module():
    spec = importlib.util.spec_from_file_location("assistant_eval_run_lsf",
                                                  _ROOT / "evals" / "assistant" / "run.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class _Res:
    def __init__(self, events, reply="", trace=None):
        self.events, self.reply, self.trace = events, reply, trace or []
        self.stopped, self.error, self.steps = "", "", 1


def test_claim_note_accepts_notes_or_a_save_note_proposal():
    mod = _eval_module()
    case = next(c for c in json.loads((_ROOT / "evals" / "assistant" / "cases.json").read_text())["cases"]
                if c["id"] == "claim_note")
    alts = case["expect"]["any_of"]
    assert {"event": "intent", "target": "notes"} in alts
    assert {"event": "confirm", "tool": "save_note"} in alts
    notes = _Res([{"type": "intent", "target": "notes"}], "I've opened Notes.")
    assert mod.evaluate(case, notes, 0)["pass"]
    proposal = _Res([{"type": "confirm", "tool": "save_note", "id": "x"}],
                    "The note is ready: confirm it above.", [{"name": "save_note"}])
    assert mod.evaluate(case, proposal, 0)["pass"]
    # Unbacked: the reply claims a saved note, but nothing was opened or proposed.
    assert not mod.evaluate(case, _Res([], "I've saved the note."), 0)["pass"]
    # A different card doesn't back it either.
    other = _Res([{"type": "confirm", "tool": "create_report", "id": "y"}],
                 "Confirm it above.", [{"name": "create_report"}])
    assert not mod.evaluate(case, other, 0)["pass"]
