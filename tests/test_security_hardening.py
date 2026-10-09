"""Security hardening: public pages without live data, role enforcement,
CSRF origin check, .env injection, formula injection, login throttling,
/backup/db consistency and a /health that leaks nothing."""
import io
import sqlite3
import time

import pytest
from conftest import RUN_ID
from openpyxl import load_workbook

ENV_USER, ENV_PASS = "envadmin", "env-password"


def _login_as(client, seeded_db, monkeypatch, username, role):
    monkeypatch.setenv("AUTH_USER", ENV_USER)
    monkeypatch.setenv("AUTH_PASS", ENV_PASS)
    client_id = seeded_db.list_clients()[0]["id"] if role == "client_viewer" else None
    seeded_db.add_user(username, f"{username}-password", role=role, client_id=client_id)
    resp = client.post("/login", data={"username": username, "password": f"{username}-password"},
                       follow_redirects=False)
    client.cookies.set("sid", resp.cookies["sid"])


# --------------------------------------------------------------------- #
# 1. Public pages: no live data for anonymous visitors when auth is on
# --------------------------------------------------------------------- #
def test_public_pages_hide_live_data_without_a_session(client, seeded_db, monkeypatch):
    seeded_db.set_setting("report_query", "Secret Subject Name")
    monkeypatch.setenv("AUTH_USER", ENV_USER)
    monkeypatch.setenv("AUTH_PASS", ENV_PASS)

    resp = client.get("/onboarding")

    assert resp.status_code == 200
    assert resp.context["live_data"] is False
    assert resp.context["stats"]["total_results"] == 0
    assert resp.context["stats"]["total_runs"] == 0
    assert "Secret Subject Name" not in resp.text
    assert resp.context["has_runs"] is False
    assert resp.context["runs_count"] == 0


def test_public_pages_show_live_data_when_auth_is_disabled(client, seeded_db):
    resp = client.get("/onboarding")
    assert resp.context["live_data"] is True
    assert resp.context["stats"]["total_results"] > 0
    assert resp.context["runs_count"] > 0


def test_public_pages_show_live_data_to_a_logged_in_user(client, seeded_db, monkeypatch):
    _login_as(client, seeded_db, monkeypatch, "alice", "operator")
    resp = client.get("/onboarding")
    assert resp.context["stats"]["total_results"] > 0
    assert resp.context["live_data"] is True


# --------------------------------------------------------------------- #
# 2. Roles
# --------------------------------------------------------------------- #
@pytest.mark.parametrize("method,path,data", [
    ("post", "/settings/apikey", {"apikey": "x"}),
    ("post", "/settings/api-keys", {"name": "n8n", "scope": "read"}),
    ("post", "/settings/api-keys/revoke", {"key_id": "1"}),
    ("post", "/settings/review-threshold", {"review_threshold": "0.5"}),
    ("get", "/backup/db", None),
    ("get", "/backup/csv", None),
    ("post", "/setup/integrations", {"gemini_key": "x"}),
    ("post", "/settings/notify", {"channel": "slack", "slack_webhook": "https://x.example/h"}),
    ("post", "/settings/cron", {"name": "weekly_report", "frequency": "daily", "hour": "3"}),
    ("get", "/audit/export.csv", None),
])
def test_operator_cannot_touch_admin_areas(client, seeded_db, monkeypatch, method, path, data):
    _login_as(client, seeded_db, monkeypatch, "alice", "operator")
    resp = getattr(client, method)(path, **({"data": data} if data else {}),
                                   follow_redirects=False)
    assert resp.status_code == 403
    assert "text/html" in resp.headers["content-type"]


def test_operator_keeps_day_to_day_writes(client, seeded_db, monkeypatch):
    _login_as(client, seeded_db, monkeypatch, "alice", "operator")
    resp = client.post("/alerts/seen", follow_redirects=False)
    assert resp.status_code != 403


def test_client_viewer_is_read_only(client, seeded_db, monkeypatch):
    _login_as(client, seeded_db, monkeypatch, "viewer", "client_viewer")
    assert client.post("/alerts/seen", follow_redirects=False).status_code == 403
    assert client.post("/settings/review-threshold", data={"review_threshold": "0.5"},
                       follow_redirects=False).status_code == 403
    assert client.get("/results").status_code == 200


_ADMIN_SECTIONS = ("<section id='search'", "<section id='ai'", "<section id='notifications'",
                   "<section id='users'", "<section id='api'", "<section id='audit'",
                   "<section id='backup'")


@pytest.mark.parametrize("role", ["operator", "client_viewer"])
def test_consolidated_settings_hide_admin_sections_from_non_admins(client, seeded_db, monkeypatch,
                                                                    role):
    _login_as(client, seeded_db, monkeypatch, "alice", role)
    resp = client.get("/settings")
    assert resp.status_code == 200
    assert "<section id='projects'" in resp.text and "<details class='subj-domains' id='domains'>" in resp.text
    for marker in _ADMIN_SECTIONS:
        assert marker not in resp.text, marker
    assert "row_hash" not in resp.text and "API keys" not in resp.text


def test_old_admin_urls_only_redirect_and_leak_nothing(client, seeded_db, monkeypatch):
    _login_as(client, seeded_db, monkeypatch, "viewer", "client_viewer")
    resp = client.get("/audit", follow_redirects=False)
    assert resp.status_code == 303 and resp.headers["location"].endswith("#audit")
    assert "<section id='audit'" not in client.get("/settings").text
    resp = client.get("/audit/export.csv", follow_redirects=False)
    assert resp.status_code == 403 and "row_hash" not in resp.text


def test_viewer_cannot_manage_subjects_or_projects(client, seeded_db, monkeypatch):
    _login_as(client, seeded_db, monkeypatch, "viewer", "client_viewer")
    assert client.post("/projects/add", data={"name": "P"}, follow_redirects=False).status_code == 403
    assert client.post("/subjects/add", data={"name": "X", "countries": ["us"]},
                       follow_redirects=False).status_code == 403
    assert "New project" not in client.get("/settings").text


def test_admin_can_read_the_audit_log(client, seeded_db, monkeypatch):
    _login_as(client, seeded_db, monkeypatch, "root", "admin")
    assert "<section id='audit'" in client.get("/audit").text
    resp = client.get("/audit/export.csv")
    assert resp.status_code == 200 and resp.text.startswith("id,at,")


def test_admin_can_manage_settings_and_backups(client, seeded_db, monkeypatch):
    _login_as(client, seeded_db, monkeypatch, "root", "admin")
    resp = client.post("/settings/api-keys", data={"name": "n8n", "scope": "read"})
    assert resp.status_code == 200
    assert seeded_db.list_api_keys()
    assert client.get("/backup/csv").status_code == 200


def test_auth_disabled_keeps_single_user_behavior(client, seeded_db):
    resp = client.post("/settings/api-keys", data={"name": "local", "scope": "read"})
    assert resp.status_code == 200
    assert client.get("/backup/csv").status_code == 200


# --------------------------------------------------------------------- #
# 3. CSRF
# --------------------------------------------------------------------- #
def test_cross_site_origin_is_rejected(client):
    resp = client.post("/settings/review-threshold", data={"review_threshold": "0.5"},
                       headers={"Origin": "http://evil.example"})
    assert resp.status_code == 403


def test_cross_site_referer_is_rejected_when_no_origin(client):
    resp = client.post("/settings/review-threshold", data={"review_threshold": "0.5"},
                       headers={"Referer": "http://evil.example/page"})
    assert resp.status_code == 403


def test_null_origin_is_rejected(client):
    resp = client.post("/login", data={"username": "a", "password": "b"},
                       headers={"Origin": "null"}, follow_redirects=False)
    assert resp.status_code == 403


def test_same_origin_and_headerless_posts_pass(client, seeded_db):
    same = client.post("/settings/review-threshold", data={"review_threshold": "0.4"},
                       headers={"Origin": "http://testserver"})
    assert same.status_code == 200
    bare = client.post("/settings/review-threshold", data={"review_threshold": "0.6"})
    assert bare.status_code == 200
    assert seeded_db.get_setting("review_threshold", "") == "0.6"


def test_api_v1_is_exempt_from_the_origin_check(client):
    resp = client.patch("/api/v1/alerts/1", json={"status": "acknowledged"},
                        headers={"Origin": "http://evil.example"})
    assert "Cross-site" not in resp.text


# --------------------------------------------------------------------- #
# 4. .env injection
# --------------------------------------------------------------------- #
@pytest.mark.parametrize("key,value", [
    ("APIFY_API_KEY", "abc\nAUTH_PASS=owned"),
    ("APIFY_API_KEY", "abc\rAUTH_PASS=owned"),
    ("APIFY_API_KEY", "abc\x00"),
    ("BAD\nKEY", "abc"),
])
def test_set_env_rejects_control_characters(tmp_path, key, value):
    from dashboard import admin

    env = tmp_path / ".env"
    env.write_text("EXISTING=1\n")
    with pytest.raises(ValueError):
        admin.set_env(str(env), key, value)
    assert env.read_text() == "EXISTING=1\n"


def test_settings_route_surfaces_invalid_env_value(client, monkeypatch, tmp_path):
    env = tmp_path / ".env"
    env.write_text("")
    monkeypatch.setenv("ENV_FILE", str(env))
    resp = client.post("/settings/email", data={"host": "smtp.example", "alert_to": "a@b\nAUTH_PASS=x"})
    assert resp.status_code == 200
    assert "line breaks" in resp.text
    assert env.read_text() == ""  # nothing half-saved


# --------------------------------------------------------------------- #
# 5. /backup/db
# --------------------------------------------------------------------- #
def test_backup_db_is_a_consistent_copy_of_the_configured_db(client, seeded_db, tmp_path):
    resp = client.get("/backup/db")
    assert resp.status_code == 200
    out = tmp_path / "copy.db"
    out.write_bytes(resp.content)
    conn = sqlite3.connect(out)
    try:
        assert conn.execute("SELECT COUNT(*) FROM search_results").fetchone()[0] == 5
    finally:
        conn.close()


# --------------------------------------------------------------------- #
# 6. Formula injection
# --------------------------------------------------------------------- #
def test_neutralize_formula():
    from export.excel_exporter import neutralize_formula

    for bad in ("=1+1", "+1", "-1", "@SUM(A1)", "\tx", "\rx"):
        assert neutralize_formula(bad) == "'" + bad
    assert neutralize_formula("Plain title") == "Plain title"
    assert neutralize_formula(3) == 3
    assert neutralize_formula(None) is None


_EVIL = {"query": "Jordan Ashworth", "url": "https://evil.example/x",
         "title": "=HYPERLINK(\"http://evil.example\",\"click\")", "description": "@cmd",
         "displayed_url": "evil.example", "position": 6, "country_code": "us",
         "domain": "evil.example", "classification": "neutral"}


def test_backup_csv_neutralizes_formulas(client, seeded_db):
    seeded_db.save_results([_EVIL], run_id=RUN_ID)
    text = client.get("/backup/csv").text
    assert "'=HYPERLINK" in text
    assert "'@cmd" in text


def test_excel_export_neutralizes_formulas(seeded_db, tmp_path):
    from export.excel_exporter import build_workbook

    seeded_db.save_results([_EVIL], run_id=RUN_ID)
    out = tmp_path / "x.xlsx"
    build_workbook(seeded_db, str(out), countries=["us"])
    wb = load_workbook(io.BytesIO(out.read_bytes()))
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert cells
    for c in cells:
        assert c.value.startswith("'=") and c.data_type != "f"


# --------------------------------------------------------------------- #
# 9. Login throttling
# --------------------------------------------------------------------- #
def test_login_is_throttled_after_repeated_failures(client, seeded_db, monkeypatch):
    import dashboard.routers.auth as auth_module

    monkeypatch.setenv("AUTH_USER", ENV_USER)
    monkeypatch.setenv("AUTH_PASS", ENV_PASS)
    seeded_db.add_user("alice", "alice-password", role="operator")

    for _ in range(auth_module._MAX_FAILURES):
        client.post("/login", data={"username": "alice", "password": "wrong"},
                    follow_redirects=False)
    locked = client.post("/login", data={"username": "alice", "password": "alice-password"},
                         follow_redirects=False)
    assert "Too+many" in locked.headers["location"]
    assert "sid" not in locked.cookies
    assert any(r["action"] == "login_locked" for r in seeded_db.list_audit(limit=20))

    # Another username from the same IP is unaffected.
    seeded_db.add_user("bob", "bob-password", role="operator")
    ok = client.post("/login", data={"username": "bob", "password": "bob-password"},
                     follow_redirects=False)
    assert "sid" in ok.cookies

    # Once the lockout expires, the right password works again.
    for entry in auth_module._FAILED_LOGINS.values():
        entry[1] = time.monotonic() - 1
    ok = client.post("/login", data={"username": "alice", "password": "alice-password"},
                     follow_redirects=False)
    assert ok.headers["location"] == "/"
    assert not auth_module._FAILED_LOGINS


# --------------------------------------------------------------------- #
# 11. /health
# --------------------------------------------------------------------- #
def test_health_does_not_leak_exception_text_or_disk_details(client, monkeypatch):
    from dashboard import deps

    def boom():
        raise RuntimeError("/secret/path/seo_dashboard.db is locked")

    monkeypatch.setattr(deps.db, "_connect", boom)
    resp = client.get("/health")
    assert resp.status_code == 503
    assert "/secret/path" not in resp.text
    assert "disk_free_gb" not in resp.json()
