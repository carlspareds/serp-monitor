"""
Excel export (redesigned)
================================
Generates an executive workbook:

  - SUM  : summary panel — KPIs, per-country table with data bars,
           charts (donut + bars) and a matrix of "own URLs per country".
  - <COUNTRY>: `Pos | URL | Title | Description | Class` table with color,
           AutoFilter, color scale on the position and frozen panes.
  - HIST : ranking by date with a DROPDOWN to pick/compare two dates
           (Pos(A), Pos(B), Δ columns with formulas), color scale and
           a line chart with an inverted Y axis.

Built on top of the DB (stored state), no network calls.
"""

import os
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Characters that make a spreadsheet app treat a cell as a formula.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def neutralize_formula(value):
    """Formula-injection guard for scraped text (titles, URLs, snippets)
    written to XLSX/CSV: a string starting with a formula character gets a
    leading single quote so it's shown as text, never evaluated. Non-strings
    pass through unchanged."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


# Palette (consistent with the manual Excel: blue=positive, red=negative, grey=neutral)
C_BLUE = "1F4E9C"
C_RED = "C62828"
C_GREY = "7F7F7F"
C_GREEN = "2E7D32"
C_DARK = "14233F"

_HEADER_FILL = PatternFill("solid", fgColor=C_DARK)
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=15, color="FFFFFF")
_SUBTITLE = Font(italic=True, color="666666", size=10)
_POS_FILL = PatternFill("solid", fgColor="E7EEFC")
_NEG_FILL = PatternFill("solid", fgColor="FDEAEA")
_THIN = Side(style="thin", color="D9DEE8")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _L(n: int) -> str:
    return get_column_letter(n)


def _style_header(ws, row: int, c1: int, c2: int, primary_hex: str = C_DARK):
    """primary_hex (5.1): client brand color -- defaults to C_DARK
    (identical to before) for any caller that does not pass it."""
    fill = _HEADER_FILL if primary_hex == C_DARK else PatternFill("solid", fgColor=primary_hex)
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _widths(ws, widths: dict[int, int]):
    for col, w in widths.items():
        ws.column_dimensions[_L(col)].width = w


# The one separator every report file (PDF, Excel, CSV sidecar) prints between
# project and subject: db.subject_label() says "Project ▸ Subject", but ▸ is
# not in the PDF font (cp1252) -- › is, and reads the same everywhere.
REPORT_SEPARATOR = "›"


def report_label(text: str | None) -> str:
    """`text` ("Project ▸ Subject") with the report separator ("Project › Subject")."""
    return (text or "").replace("▸", REPORT_SEPARATOR).strip()


def subject_title(db, subject_id: int | None = None) -> str:
    """Title of a report of the subject `db` is scoped to (or `subject_id`):
    "Project › Subject" (db.subject_label), else the subject's name / tracked
    query, else the SUBJECT_QUERY env var (CLI scripts on an empty DB). Every
    report route (Reports dialog, /export/*, weekly job, assistant) goes
    through here, so they all print the same label."""
    try:
        title = db.subject_label(subject_id)
    except Exception:  # noqa: BLE001 -- a title must never break an export
        title = ""
    if not title:
        try:
            subject = db.current_subject()
            title = (subject.get("name") or subject.get("query") or "").strip()
        except Exception:  # noqa: BLE001
            title = ""
    return report_label(title) or os.getenv("SUBJECT_QUERY", "")


def build_workbook(db, output_path: str, countries=None, subject_name: str | None = None) -> str:
    """Workbook of the subject `db` is scoped to. `subject_name` overrides the
    title (default subject_title(db): "Project › Subject")."""
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # Client brand (5.1): parameterized header color -- without a
    # `branding` row (first boot), get_branding() returns the same
    # C_DARK ("#14233F") as always, so the Excel comes out THE SAME as
    # before this feature. openpyxl wants the hex WITHOUT '#'.
    primary_hex = db.get_branding()["primary_color"].lstrip("#").upper()

    latest = db.latest_run_per_country()
    if countries:
        wanted = {c.lower() for c in countries}
        latest = [r for r in latest if (r.get("country_code") or "").lower() in wanted]

    # Canonical order of country sheets: GLOBAL, CA, US, HK, GB, FR, CH, DE, IT, NL
    ORDER = ["global", "ca", "us", "hk", "gb", "fr", "ch", "de", "it", "nl"]
    latest.sort(key=lambda r: ORDER.index((r.get("country_code") or "").lower())
                if (r.get("country_code") or "").lower() in ORDER else 99)

    # Results of the latest run per country -- get_results_filtered (2.3)
    # excludes same_person='no' by default ("Stats and Excel exclude them"):
    # sort_by=position/order=asc + a high limit reproduces the order of
    # get_results_by_run (a single run -> country_code is constant, only
    # position matters) without using that raw read. /run/{id}, timeline and
    # analysis DO use get_results_by_run unfiltered (drill-down of a
    # specific run); report/weekly.py and processing/alerts.py, on the
    # other hand, ALSO filter (H16/H17) because they feed the weekly email
    # and the alerts, which are in the same scope as Excel per the deep
    # dive §2.3.
    country_results = {
        (r.get("country_code") or "").lower(): db.get_results_filtered(
            run_id=r.get("apify_run_id"), sort_by="position", order="asc", limit=100000,
        )
        for r in latest
    }

    # Tab order: SUM, HIST, <countries in canonical order>
    subject = report_label(subject_name) or subject_title(db)
    _build_sum_sheet(wb, latest, country_results, primary_hex, subject=subject)
    _build_hist_sheet(wb, db.get_history_pivot(), primary_hex)
    for run in latest:
        country = (run.get("country_code") or "??").lower()
        _build_country_sheet(wb, country.upper()[:28], run, country_results.get(country, []),
                             primary_hex)

    wb.save(output_path)
    return output_path



# --------------------------------------------------------------------- #
# SUM
# --------------------------------------------------------------------- #
def _build_sum_sheet(wb: Workbook, latest: list[dict], country_results: dict[str, list[dict]],
                     primary_hex: str = C_DARK, subject: str = ""):
    ws = wb.active
    ws.title = "SUM"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    ws["A1"] = f"SERP MONITOR  ·  {subject}" if subject else "SERP MONITOR"
    ws["A1"].fill = PatternFill("solid", fgColor=primary_hex)
    ws["A1"].font = _TITLE_FONT
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:H2")
    ws["A2"] = f"Generated: {datetime.now(UTC).replace(tzinfo=None).isoformat(timespec='seconds')} UTC"
    ws["A2"].font = _SUBTITLE

    # Totals
    tot = {"positive": 0, "negative": 0, "neutral": 0, "total": 0}
    for r in latest:
        for k, key in (("positive", "positive"), ("negative", "negative"),
                       ("neutral", "neutral"), ("total", "total_results")):
            tot[k] += r.get(key) or 0
    pct_pos = round(100 * tot["positive"] / tot["total"]) if tot["total"] else 0

    # --- KPI cards (row 4 label, row 5 number) ---
    kpis = [
        ("POSITIVE", tot["positive"], C_BLUE),
        ("NEGATIVE", tot["negative"], C_RED),
        ("NEUTRAL", tot["neutral"], C_GREY),
        ("% POSITIVE", f"{pct_pos}%", C_GREEN),
        ("COUNTRIES", len(latest), primary_hex),
    ]
    col = 1
    for label, value, color in kpis:
        c1, c2 = col, col + 1
        ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
        lab = ws.cell(row=4, column=c1, value=label)
        lab.fill = PatternFill("solid", fgColor=color)
        lab.font = Font(color="FFFFFF", bold=True, size=9)
        lab.alignment = Alignment(horizontal="center")
        num = ws.cell(row=5, column=c1, value=value)
        num.fill = PatternFill("solid", fgColor="F2F5FA")
        num.font = Font(color=color, bold=True, size=20)
        num.alignment = Alignment(horizontal="center")
        ws.row_dimensions[5].height = 34
        col += 2

    # --- Per-country table ---
    # Engine dimension (1.1): "Engine" column ONLY if there is more than one
    # engine among the `latest` runs -- with just one (today's case) the
    # table stays IDENTICAL to before this item. `latest` is still "the
    # most recent run per country" (latest_run_per_country was not touched
    # in 1.1 -- see # ponytail below): with several engines per country,
    # this row represents the LAST engine saved for that country, not an
    # aggregate of all of them -- a known MVP limitation.
    # ponytail: SUM only shows, per country, the run of the LAST engine
    # saved (latest_run_per_country groups by country_code, not by
    # (country_code, engine)) -- a country with google+bing does not sum
    # both in this table. Trigger: if a future item needs the real
    # multi-engine aggregate in SUM, change latest_run_per_country to group
    # by (country_code, engine) and sum here -- out of scope for MVP 1.1
    # (avoids breaking the OTHER consumers of latest_run_per_country:
    # weekly report, PDF, ai_assistant, all of which assume one
    # row/country).
    multi_engine_sum = len({(r.get("engine") or "google") for r in latest}) > 1
    hr = 7
    headers = ["Country", "Positive", "Negative", "Neutral", "Total", "% Positive"]
    if multi_engine_sum:
        headers = [*headers, "Engine"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=hr, column=i, value=h)
    _style_header(ws, hr, 1, len(headers), primary_hex)
    r = hr + 1
    for run in latest:
        total = run.get("total_results") or 0
        pos = run.get("positive") or 0
        ws.cell(row=r, column=1, value=(run.get("country_code") or "").upper())
        ws.cell(row=r, column=2, value=pos)
        ws.cell(row=r, column=3, value=run.get("negative") or 0)
        ws.cell(row=r, column=4, value=run.get("neutral") or 0)
        ws.cell(row=r, column=5, value=total)
        ws.cell(row=r, column=6, value=round(100 * pos / total) if total else 0)
        if multi_engine_sum:
            ws.cell(row=r, column=7, value=(run.get("engine") or "google").upper())
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = _BORDER
        r += 1
    last_country_row = r - 1
    # total row
    ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=r, column=2, value=tot["positive"]).font = Font(bold=True)
    ws.cell(row=r, column=3, value=tot["negative"]).font = Font(bold=True)
    ws.cell(row=r, column=4, value=tot["neutral"]).font = Font(bold=True)
    ws.cell(row=r, column=5, value=tot["total"]).font = Font(bold=True)
    ws.cell(row=r, column=6, value=pct_pos).font = Font(bold=True)
    total_row = r

    # Data bars on "% Positive" (col 6) and color scales on Pos/Neg
    if last_country_row >= hr + 1:
        rng = f"F{hr+1}:F{last_country_row}"
        ws.conditional_formatting.add(
            rng, DataBarRule(start_type="num", start_value=0, end_type="num",
                             end_value=100, color=C_BLUE))
        ws.conditional_formatting.add(
            f"B{hr+1}:B{last_country_row}",
            ColorScaleRule(start_type="min", start_color="FFFFFF",
                           end_type="max", end_color="BBD3F0"))
        ws.conditional_formatting.add(
            f"C{hr+1}:C{last_country_row}",
            ColorScaleRule(start_type="min", start_color="FFFFFF",
                           end_type="max", end_color="F3B6B6"))

    sum_widths = {1: 10, 2: 11, 3: 11, 4: 11, 5: 9, 6: 12}
    if multi_engine_sum:
        sum_widths[7] = 10
    _widths(ws, sum_widths)

    # --- Charts: totals donut + per-country bars ---
    if latest:
        ws["J7"] = "Class"; ws["K7"] = "n"
        ws["J8"] = "Positive"; ws["K8"] = tot["positive"]
        ws["J9"] = "Negative"; ws["K9"] = tot["negative"]
        ws["J10"] = "Neutral"; ws["K10"] = tot["neutral"]
        pie = PieChart(); pie.title = "Global classification"
        pie.height, pie.width = 7.5, 10
        pie.add_data(Reference(ws, min_col=11, min_row=7, max_row=10), titles_from_data=True)
        pie.set_categories(Reference(ws, min_col=10, min_row=8, max_row=10))
        try:
            pie.holeSize = 55
        except Exception:  # noqa: BLE001
            pass
        ws.add_chart(pie, "H4")

        bar = BarChart(); bar.type = "col"; bar.title = "Classification by country"
        bar.height, bar.width = 7.5, 14
        data = Reference(ws, min_col=2, max_col=4, min_row=hr, max_row=last_country_row)
        cats = Reference(ws, min_col=1, min_row=hr + 1, max_row=last_country_row)
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(cats)
        ws.add_chart(bar, f"H{total_row + 3}")

    # --- Matrix: own (positive) URLs by country ---
    _build_own_urls_matrix(ws, latest, country_results, start_row=total_row + 22,
                           primary_hex=primary_hex)


def _build_own_urls_matrix(ws, latest, country_results, start_row, primary_hex: str = C_DARK):
    countries = [(r.get("country_code") or "").lower() for r in latest]
    matrix: dict[str, dict[str, int]] = {}
    for c in countries:
        for res in country_results.get(c, []):
            if res.get("classification") != "positive":
                continue
            url = res.get("url") or ""
            if not url:
                continue
            matrix.setdefault(url, {})[c] = res.get("position")

    ws.cell(row=start_row, column=1, value="OWN URLs (positive) — position by country")
    ws.cell(row=start_row, column=1).font = Font(bold=True, size=12, color=C_BLUE)
    hr = start_row + 1
    headers = ["URL"] + [c.upper() for c in countries]
    for i, h in enumerate(headers, 1):
        ws.cell(row=hr, column=i, value=h)
    _style_header(ws, hr, 1, len(headers), primary_hex)
    r = hr + 1
    for url, by_c in sorted(matrix.items(), key=lambda kv: -len(kv[1])):
        ws.cell(row=r, column=1, value=neutralize_formula(url))
        for j, c in enumerate(countries, 2):
            v = by_c.get(c)
            cell = ws.cell(row=r, column=j, value=("" if v is None else v))
            cell.alignment = Alignment(horizontal="center")
            if v is not None:
                cell.fill = _POS_FILL
        r += 1
    if r > hr + 1:
        ws.conditional_formatting.add(
            f"{_L(2)}{hr+1}:{_L(len(headers))}{r-1}",
            ColorScaleRule(start_type="num", start_value=1, start_color="63BE7B",
                           mid_type="num", mid_value=5, mid_color="FFEB84",
                           end_type="num", end_value=10, end_color="F8696B"))
    ws.column_dimensions["A"].width = 60


# --------------------------------------------------------------------- #
# Per-country sheet
# --------------------------------------------------------------------- #
def _build_country_sheet(wb: Workbook, country: str, run: dict, results: list[dict],
                         primary_hex: str = C_DARK):
    ws = wb.create_sheet(title=country)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:E1")
    ws["A1"] = f"Search Settings: {country}"
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=primary_hex)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)

    ws["A3"] = "Platform URLs Found"; ws["B3"] = run.get("positive") or 0
    ws["A4"] = "Negative URLs Found"; ws["B4"] = run.get("negative") or 0
    ws["A3"].font = ws["A4"].font = Font(bold=True)

    headers = ["Pos", "URL", "Title", "Description", "Class"]
    hr = 6
    for i, h in enumerate(headers, 1):
        ws.cell(row=hr, column=i, value=h)
    _style_header(ws, hr, 1, len(headers), primary_hex)
    r = hr + 1
    for res in results:
        ws.cell(row=r, column=1, value=res.get("position"))
        ws.cell(row=r, column=2, value=neutralize_formula(res.get("url")))
        ws.cell(row=r, column=3, value=neutralize_formula(res.get("title")))
        desc = (res.get("description") or "")[:160]
        ws.cell(row=r, column=4, value=neutralize_formula(desc))
        cls = res.get("classification") or "neutral"
        ws.cell(row=r, column=5, value=cls)
        fill = _POS_FILL if cls == "positive" else (_NEG_FILL if cls == "negative" else None)
        if fill:
            for c in range(1, 6):
                ws.cell(row=r, column=c).fill = fill
        r += 1
    last = r - 1
    if last >= hr + 1:
        ws.auto_filter.ref = f"A{hr}:E{last}"
        ws.conditional_formatting.add(
            f"A{hr+1}:A{last}",
            ColorScaleRule(start_type="num", start_value=1, start_color="63BE7B",
                           mid_type="num", mid_value=5, mid_color="FFEB84",
                           end_type="num", end_value=10, end_color="F8696B"))
    ws.freeze_panes = f"A{hr+1}"
    _widths(ws, {1: 6, 2: 50, 3: 42, 4: 60, 5: 12})


# --------------------------------------------------------------------- #
# HIST with comparison dropdown
# --------------------------------------------------------------------- #
def _build_hist_sheet(wb: Workbook, pivot: dict, primary_hex: str = C_DARK):
    ws = wb.create_sheet(title="HIST")
    ws.sheet_view.showGridLines = False
    dates = pivot.get("dates", [])
    # HIST only shows positive and negative results (no neutrals)
    rows = [r for r in pivot.get("rows", [])
            if r.get("classification") in ("positive", "negative")]
    # Engine dimension (1.1): Database.get_history_pivot ONLY adds "engine"
    # to each row if there is more than one engine in the data -- with just
    # one (today's case) `has_engine` is False and the sheet stays
    # IDENTICAL to before this item (same column layout).
    has_engine = any("engine" in r for r in rows)

    ws.merge_cells("A1:E1")
    ws["A1"] = "RANKING HISTORY (1 = best; blank = not present)"
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=primary_hex)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)

    # Column layout -- 1.1 inserts "Engine" in col 4 (shifting the dates
    # one position) ONLY if has_engine.
    C_URL, C_TITLE, C_CLS = 1, 2, 3
    C_ENGINE = 4 if has_engine else None
    first_date_col = 5 if has_engine else 4
    last_date_col = first_date_col + len(dates) - 1 if dates else first_date_col - 1
    col_posA = max(first_date_col, last_date_col + 1)
    col_posB = col_posA + 1
    col_delta = col_posB + 1

    # Comparison controls (row 3): A and B dropdowns
    ws["A3"] = "Compare dates:"; ws["A3"].font = Font(bold=True)
    ws["B3"] = "A ="
    ws["D3"] = "B ="
    latest_a = dates[-1] if dates else ""
    latest_b = dates[-2] if len(dates) >= 2 else (dates[-1] if dates else "")
    ws["C3"] = latest_a
    ws["E3"] = latest_b
    for cell in ("C3", "E3"):
        ws[cell].fill = PatternFill("solid", fgColor="FFF3CD")
        ws[cell].font = Font(bold=True)
    if dates:
        dv = DataValidation(type="list", formula1='"' + ",".join(dates) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(ws["C3"]); dv.add(ws["E3"])

    # Header (row 5)
    hr = 5
    ws.cell(row=hr, column=C_URL, value="URL")
    ws.cell(row=hr, column=C_TITLE, value="Title")
    ws.cell(row=hr, column=C_CLS, value="Class")
    if has_engine:
        ws.cell(row=hr, column=C_ENGINE, value="Engine")
    for j, d in enumerate(dates):
        ws.cell(row=hr, column=first_date_col + j, value=d)
    ws.cell(row=hr, column=col_posA, value="Pos (A)")
    ws.cell(row=hr, column=col_posB, value="Pos (B)")
    ws.cell(row=hr, column=col_delta, value="Δ (A→B)")
    _style_header(ws, hr, 1, col_delta, primary_hex)

    r = hr + 1
    LA, LL = _L(first_date_col), _L(last_date_col if dates else first_date_col)
    for row in rows:
        ws.cell(row=r, column=C_URL, value=neutralize_formula(row.get("url")))
        ws.cell(row=r, column=C_TITLE, value=neutralize_formula(row.get("title")))
        ws.cell(row=r, column=C_CLS, value=row.get("classification"))
        if has_engine:
            ws.cell(row=r, column=C_ENGINE, value=row.get("engine", ""))
        for j, d in enumerate(dates):
            val = row.get("positions", {}).get(d)
            ws.cell(row=r, column=first_date_col + j, value=val)  # number or None
        if dates:
            ws.cell(row=r, column=col_posA).value = (
                f"=IFERROR(INDEX(${LA}{r}:${LL}{r},MATCH($C$3,${LA}${hr}:${LL}${hr},0)),\"\")"
            )
            ws.cell(row=r, column=col_posB).value = (
                f"=IFERROR(INDEX(${LA}{r}:${LL}{r},MATCH($E$3,${LA}${hr}:${LL}${hr},0)),\"\")"
            )
            pa, pb = f"{_L(col_posA)}{r}", f"{_L(col_posB)}{r}"
            ws.cell(row=r, column=col_delta).value = (
                f"=IF(AND(ISNUMBER({pa}),ISNUMBER({pb})),{pa}-{pb},\"\")"
            )
        r += 1
    last_row = r - 1

    # Color scale on date columns and on Pos(A)/Pos(B)
    if dates and last_row >= hr + 1:
        def cs():
            return ColorScaleRule(start_type="num", start_value=1, start_color="63BE7B",
                                  mid_type="num", mid_value=5, mid_color="FFEB84",
                                  end_type="num", end_value=10, end_color="F8696B")
        ws.conditional_formatting.add(f"{LA}{hr+1}:{LL}{last_row}", cs())
        ws.conditional_formatting.add(f"{_L(col_posA)}{hr+1}:{_L(col_posB)}{last_row}", cs())

    # Line chart (inverted Y axis), up to 10 URLs
    if dates and last_row >= hr + 1:
        chart = LineChart()
        chart.title = "Ranking evolution by URL"
        chart.height, chart.width = 9, 20
        chart.y_axis.title = "Position (1 = best)"
        chart.y_axis.scaling.orientation = "maxMin"
        max_series_row = min(last_row, hr + 10)
        data = Reference(ws, min_col=C_TITLE, max_col=last_date_col,
                         min_row=hr + 1, max_row=max_series_row)
        cats = Reference(ws, min_col=first_date_col, max_col=last_date_col, min_row=hr)
        chart.add_data(data, titles_from_data=True, from_rows=True)
        chart.set_categories(cats)
        ws.add_chart(chart, f"A{last_row + 3}")

    ws.freeze_panes = f"{_L(first_date_col)}{hr+1}"
    widths = {1: 50, 2: 42, 3: 10}
    if has_engine:
        widths[C_ENGINE] = 10
    for j in range(first_date_col, col_delta + 1):
        widths[j] = 11
    _widths(ws, widths)
